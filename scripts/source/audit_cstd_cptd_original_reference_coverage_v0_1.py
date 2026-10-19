#!/usr/bin/env python3
"""Outcome-blind source-lineage coverage of CSTD × CPTDv2 organ-pair candidates.

Only taxon names, trait-label equivalence (unverified candidates), presence,
site/source citation identifiers and frozen original pair identities are used.
Never aggregate numerical phenotypes, phylogenetic K or effect directions.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict, Counter
from pathlib import Path

from openpyxl import load_workbook

from audit_cstd_cptd_pair_support_v0_1 import (
    EXPECTED_MD5, CSTD_TRAITS, CPTD_TRAITS, REPS, VEGS,
    tokenized_binomial, from_translation, present, verify
)

SITE_MD5 = "76debd1ec32b850efa2c593ee5b94ee4"


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def evidence_spans(sources_by_species, species):
    """Sum source support at distinct SPECIES resolution, never rows or trait-pairs."""
    out = defaultdict(set)
    for name in species:
        for source in sources_by_species.get(name, set()):
            if source:
                out[source].add(name)
    return {s: len(names) for s, names in out.items()}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--cstd", required=True, type=Path)
    p.add_argument("--translation", required=True, type=Path)
    p.add_argument("--chemical", required=True, type=Path)
    p.add_argument("--sites", required=True, type=Path)
    p.add_argument("--source-support", required=True, type=Path)
    p.add_argument("--cstd-support", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    args = p.parse_args()
    for key, path in (("cstd", args.cstd), ("translation", args.translation),
                      ("chemical", args.chemical)):
        verify(path, EXPECTED_MD5[key])
    verify(args.sites, SITE_MD5)
    locked = json.loads(args.source_support.read_text(encoding="utf-8"))
    original_cstd = json.loads(args.cstd_support.read_text(encoding="utf-8"))
    assert locked["status"] == "HOLD_PRE_OUTCOME"
    assert original_cstd["status"] == "HOLD_PRE_OUTCOME"

    wb = load_workbook(args.cstd, read_only=True, data_only=True)
    ws = wb["Seed trait data"]
    cols = {str(c):i for i,c in enumerate(next(ws.iter_rows(values_only=True)))}
    assert {"Family_accepted","Species_accepted","Trait","Reference","Value"}.issubset(cols)
    cstd_sources = defaultdict(set)
    # Reproduce the CSTD-ONLY frozen pair roster from exact original
    # Species_accepted strings, NEVER from taxonomic normalization.
    cstd_sources_raw = defaultdict(set)
    canonical_fam = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Value presence is the only read permitted here. No numeric conversion.
        if not present(row[cols["Value"]]):
            continue
        trait = CSTD_TRAITS.get(norm(row[cols["Trait"]]))
        family = str(row[cols["Family_accepted"]] or "").strip()
        original_species = str(row[cols["Species_accepted"]] or "").strip()
        species = tokenized_binomial(original_species)
        ref = str(row[cols["Reference"]] or "").strip()
        if not trait or not family or not species or not ref:
            continue
        cstd_sources[(family, species, trait)].add(ref)
        cstd_sources_raw[(family, original_species, trait)].add(ref)
        canonical_fam[species].add(family)
    conflicts = {sp for sp, fs in canonical_fam.items() if len(fs) != 1}
    for key in list(cstd_sources):
        if key[1] in conflicts:
            del cstd_sources[key]

    with args.sites.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        names = reader.fieldnames or []
        idkey = next((n for n in names if norm(n) == "siteid"), None)
        src_key = next((n for n in names if norm(n) == "source"), None)
        samp_key = next((n for n in names if norm(n) == "samplingstrategy"), None)
        if not idkey or not src_key:
            raise RuntimeError("Site ID and source labels unavailable; fail closed")
        sites = {}
        for row in reader:
            site_id = str(row[idkey] or "").strip()
            if site_id:
                sites[site_id] = {
                    "source": str(row[src_key] or "").strip(),
                    "strategy": str(row[samp_key] or "").strip() if samp_key else "",
                }

    with args.translation.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        assert {"SAMPLE ID","Site ID","ACCEPTED GENUS","ACCEPTED SPECIES"}.issubset(reader.fieldnames or [])
        samples = {}
        for row in reader:
            species = from_translation(row)
            if species is not None:
                samples[str(row["SAMPLE ID"]).strip()] = (
                    species, str(row["Site ID"] or "").strip())

    leaf_species_sites = defaultdict(set)
    leaf_species_source = defaultdict(set)
    leaf_species_strategy = defaultdict(set)
    present_samples = 0
    unresolved_sites = 0
    with args.chemical.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        assert {"SAMPLE ID","Average LA","LMA"}.issubset(reader.fieldnames or [])
        for row in reader:
            key = str(row["SAMPLE ID"]).strip()
            if key not in samples:
                continue
            species, site_id = samples[key]
            for column, trait in CPTD_TRAITS.items():
                if not present(row[column]):
                    continue
                present_samples += 1
                leaf_species_sites[(species, trait)].add(site_id)
                if site_id in sites:
                    info = sites[site_id]
                    if info["source"]:
                        leaf_species_source[(species, trait)].add(info["source"])
                    if info["strategy"]:
                        leaf_species_strategy[(species, trait)].add(info["strategy"])
                else:
                    unresolved_sites += 1

    pair_data = locked["metadata_support"]["cross_source_pairs"]
    assert len(pair_data) == locked["metadata_support"]["cross_source_20plus_pairs"]
    assert len(pair_data) == 17, "Frozen pair identity changed unexpectedly; HOLD"
    cross_family = defaultdict(lambda: {"pairs": [],"species": set(),"rep_sources": defaultdict(set),"leaf_sources":defaultdict(set),"strategies":set()})
    for x in pair_data:
        family = x["family"]
        rep = x["reproductive"]
        leaf = x["vegetative"]
        expected_n = x["same_accepted_binomial"]
        focal = {
            sp for (fam,sp,trait) in cstd_sources
            if fam==family and trait==rep and (sp,leaf) in leaf_species_sites
        }
        if len(focal) != expected_n:
            raise RuntimeError("Frozen matched-species identity mismatch; HOLD")
        rec = cross_family[family]
        rec["pairs"].append((rep, leaf, expected_n))
        rec["species"].update(focal)
        for sp in focal:
            rec["rep_sources"][sp].update(cstd_sources[(family,sp,rep)])
            rec["leaf_sources"][sp].update(leaf_species_source.get((sp,leaf),set()))
            rec["strategies"].update(leaf_species_strategy.get((sp,leaf),set()))

    family_reports = []
    all_rep_refs = set()
    all_leaf_refs = set()
    for family, dat in sorted(cross_family.items()):
        sp = dat["species"]
        rep_ref = evidence_spans(dat["rep_sources"],sp)
        leaf_ref = evidence_spans(dat["leaf_sources"],sp)
        all_rep_refs.update(rep_ref)
        all_leaf_refs.update(leaf_ref)
        family_reports.append({
            "family":family,
            "pair_names_and_accepted_binomial_support": [
                {"reproductive":r,"vegetative":v,"n":n} for r,v,n in dat["pairs"]
            ],
            "unique_species_any_qualifying_leaf_pair":len(sp),
            "fraction_with_cstd_reproductive_source_labels":
                sum(bool(dat["rep_sources"].get(x)) for x in sp)/len(sp),
            "fraction_with_site_source_labels":
                sum(bool(dat["leaf_sources"].get(x)) for x in sp)/len(sp),
            "distinct_cstd_original_reference_labels":len(rep_ref),
            "distinct_cptd_site_source_labels":len(leaf_ref),
            "max_species_fraction_per_single_cstd_reference":
                round(max(rep_ref.values(),default=0)/len(sp),4),
            "max_species_fraction_per_single_cptd_source":
                round(max(leaf_ref.values(),default=0)/len(sp),4),
            "cptd_sampling_strategy_metadata":sorted(dat["strategies"]),
            "same_original_individuals_between_databases":"NOT_ESTABLISHED",
            "verified_independent_primary_source":"UNKNOWN"
        })

    # For CSTD-only organ pairs, use the EXACT ORIGINAL accepted-name string
    # just as the immutable CSTD v0.1 support extractor did. Normalized
    # binomials can COMBINE distinct raw name rows with complementary traits
    # and artificially CREATE supported pairs. Never do that in this arm.
    # Site-crosswalk candidate counts above remain separate, provisional.
    same_ref_pairs = []
    for item in original_cstd["metadata_support_only"].get("pairs",[]):
        family, rep, veg = item["family"],item["reproductive"],item["vegetative"]
        if veg!="plant_height":
            continue
        species = {
            sp for (fam,sp,trait) in cstd_sources_raw
            if fam==family and trait==rep and (fam,sp,veg) in cstd_sources_raw
        }
        if len(species) != item["same_species"]:
            raise RuntimeError("Exact raw CSTD species support failed identity check; HOLD")
        # Report stricter family-unambiguous support as a second diagnostic
        # without changing the frozen original cohort.
        stricter = {
            sp for sp in species
            if tokenized_binomial(sp) not in conflicts
        }
        removed_ambiguous = len(species)-len(stricter)
        within_source_sp = {
            sp for sp in stricter
            if bool(cstd_sources_raw[(family,sp,rep)] & cstd_sources_raw[(family,sp,veg)])
        }
        same_ref_pairs.append({
            "family":family,"rep":rep,"vegetative":veg,
            "same_species":len(stricter),
            "original_same_species_before_conflict_exclusion":len(species),
            "removed_conflicting_family_identity":removed_ambiguous,
            "same_cited_reference_species":len(within_source_sp),
            "at_least_20_same_reference_species":len(within_source_sp)>=20
        })

    out = {
        "status":"HOLD_PRE_OUTCOME",
        "numeric_trait_values_used":False,
        "K_estimated":False,
        "source_identity_md5_verified":True,
        "source_observation_independence":"NOT_ESTABLISHED",
        "source_provenance_level":"CSTD cited reference and CPTD site source, NOT raw specimen/field record",
        "cross_source": {
            "families_with_20plus_leaf_pair":len(family_reports),
            "pairs":len(pair_data),
            "families":family_reports,
            "distinct_cstd_reference_labels_across_eligible_families":len(all_rep_refs),
            "distinct_cptd_site_source_labels_across_eligible_families":len(all_leaf_refs),
            "cptd_site_count":len(sites),
            "cptd_leaf_sample_trait_presence_rows":present_samples,
            "leaf_trait_presences_with_unresolved_site":unresolved_sites
        },
        "cstd_within_original_reference": {
            "input_CSTD_only_pairs":len(same_ref_pairs),
            "pairs_retaining_20plus_species_with_cited_reference_shared_by_both_traits":
                sum(q["at_least_20_same_reference_species"] for q in same_ref_pairs),
            "families_retaining_at_least_one_20plus_same_reference_pair":
                len({q["family"] for q in same_ref_pairs if q["at_least_20_same_reference_species"]}),
            "pairs":same_ref_pairs,
        },
        "hard_stop": "Cited source ID overlap is not same individual or same direct primary measurement. Original bibliographic lineage, AusTraits reuse and two phylogenies remain unresolved. No K."
    }
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(out,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print("CSTD_CPTD_PROVENANCE_COVERAGE="+json.dumps({
        "status":out["status"],
        "cross_source_families":out["cross_source"]["families_with_20plus_leaf_pair"],
        "family_summary":[{k:v for k,v in row.items() if k in (
            "family","unique_species_any_qualifying_leaf_pair",
            "distinct_cstd_original_reference_labels","distinct_cptd_site_source_labels",
            "max_species_fraction_per_single_cstd_reference",
            "max_species_fraction_per_single_cptd_source",
            "fraction_with_site_source_labels")} for row in family_reports],
        "within_reference_20plus_pairs":out["cstd_within_original_reference"]["pairs_retaining_20plus_species_with_cited_reference_shared_by_both_traits"],
        "within_reference_20plus_families":out["cstd_within_original_reference"]["families_retaining_at_least_one_20plus_same_reference_pair"],
    },ensure_ascii=False))


if __name__=="__main__":
    main()
