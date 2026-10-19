#!/usr/bin/env python3
"""Outcome-blind CSTD fruit/seed x NHM vegetative trait source-support screen.

Input source identities pinned to CSTD Figshare v4 MD5 and exact NHM Git commit
plus SHA256 for main files. Only measurement PRESENCE is used. Never parse,
transform, summarize or compare phenotype magnitudes or phylogenetic K.

NHM floras are SECONDARY BOTANICAL DESCRIPTIONS and do NOT qualify as new
independent primary measurements without record-level provenance adjudication.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

from audit_nhm_organ_presence_v0_1 import (
    SHA256, SOURCE_SHA, SOURCES, REPRODUCTIVE, VEGETATIVE, CV_VEGETATIVE,
    binomial, family_map, source_presence, cv_presence, check, is_present
)

CSTD_MD5 = "9d02dfed7de29484d2f5c5137235cabe"
CSTD_ALLOWED_REPRODUCTIVE = {
    "fruitlength": "fruit_length", "fruitwidth": "fruit_width",
    "seedlength": "seed_length", "seedwidth": "seed_width"
}
CSTD_SOURCE_FIELDS = ("Species_accepted", "Family_accepted", "Trait", "Value", "Reference")

def norm(x):
    return re.sub(r"[^a-z0-9]", "", str(x or "").lower())

def cstd_presence(path, wfomap):
    h = hashlib.md5()
    with path.open("rb") as fh:
        for block in iter(lambda:fh.read(1024*1024), b""):
            h.update(block)
    if h.hexdigest() != CSTD_MD5:
        raise RuntimeError("Unexpected CSTD workbook input identity")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Seed trait data"]
    header = {str(v):i for i,v in enumerate(next(ws.iter_rows(values_only=True)))}
    if not set(CSTD_SOURCE_FIELDS).issubset(header):
        raise RuntimeError("CSTD schema changed, fail closed")
    cstd_states = defaultdict(set)
    cstd_families = defaultdict(set)
    cstd_sources = defaultdict(set)
    rows = 0
    eligible_rows = 0
    raw_species_names = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows += 1
        trait = CSTD_ALLOWED_REPRODUCTIVE.get(norm(row[header["Trait"]]))
        if trait is None:
            continue
        # No conversion of numeric trait values. Presence flag only.
        if not is_present(row[header["Value"]]):
            continue
        raw_name = str(row[header["Species_accepted"]] or "").strip()
        parts = re.findall("[A-Za-z-]+", raw_name)
        if len(parts) < 2:
            continue
        candidate = binomial(parts[0] + " " + parts[1])
        family = str(row[header["Family_accepted"]] or "").strip()
        reference = str(row[header["Reference"]] or "").strip()
        if candidate is None or candidate not in wfomap or not family or not reference:
            continue
        cstd_families[candidate].add(family)
        raw_species_names[candidate].add(raw_name)
        if wfomap[candidate] != family:
            continue
        cstd_states[candidate].add(trait)
        cstd_sources[(candidate,trait)].add(reference)
        eligible_rows+=1
    conflicted = {sp for sp,fs in cstd_families.items() if len(fs)>1}
    for sp in conflicted:
        cstd_states.pop(sp,None)
    return cstd_states, {
        "input_rows":rows, "WFO_accepted_family_matched_rows":eligible_rows,
        "WFO_accepted_family_matched_species":len(cstd_states),
        "conflicting_CSTD_family_species_removed":len(conflicted),
        "mapped_binomial_with_multiple_original_texts":
            sum(len(names)>1 for sp,names in raw_species_names.items() if sp in cstd_states),
        "reference_labels_in_WFO_matched_data":
            len({ref for (sp,trait),refs in cstd_sources.items()
                 if sp in cstd_states for ref in refs}),
    }

def combined_support(repro_states, veg_states, family, source_name, min_species=20):
    pair_species=defaultdict(set)
    for sp,reps in repro_states.items():
        vegs=veg_states.get(sp,set())
        for rep in sorted(reps):
            for veg in sorted(vegs):
                if veg not in VEGETATIVE and veg not in CV_VEGETATIVE:
                    continue
                pair_species[(family[sp], rep, veg)].add(sp)
    eligible=[{"family":fa,"reproductive":rep,"vegetative":veg,
               "same_species":len(names)}
               for (fa,rep,veg),names in sorted(pair_species.items())
               if len(names)>=min_species]
    return {
        "source":source_name,
        "same_taxon_20plus_family_trait_pair_count":len(eligible),
        "family_count":len({z["family"] for z in eligible}),
        "families_fruit":len({z["family"] for z in eligible
                              if z["reproductive"].startswith("fruit_")}),
        "families_seed":len({z["family"] for z in eligible
                             if z["reproductive"].startswith("seed_")}),
        "reproductive_trait_names":sorted({z["reproductive"] for z in eligible}),
        "vegetative_trait_names":sorted({z["vegetative"] for z in eligible}),
        "family_names":sorted({z["family"] for z in eligible}),
        "pair_support":eligible,
        "candidate_not_independent_primary_validation":True,
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--cstd",type=Path,required=True)
    p.add_argument("--wfo",type=Path,required=True)
    p.add_argument("--cv",type=Path,required=True)
    for source in SOURCES:
        p.add_argument("--"+source.replace("_","-"),type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    a=p.parse_args()
    check(a.wfo,SHA256["wfo"])
    check(a.cv,SHA256["cv"])
    f,conflict,invalid=family_map(a.wfo)
    cstd, cstd_audit=cstd_presence(a.cstd,f)
    result = {
        "status":"HOLD_PRE_OUTCOME_PRIMARY_PROVENANCE",
        "sources":["CSTD_v4", "NaturalHistoryMuseum/plant-traits"],
        "nhm_source_sha":SOURCE_SHA,
        "cstd_source_md5":CSTD_MD5,
        "meta_only":True,
        "numeric_magnitudes_analyzed":False,
        "K_evaluated":False,
        "pair_grain":"one family × one CSTD reproductive trait × one NHM vegetative trait; exactly matched accepted two-token binomial supported by WFO family",
        "wfo_excluded_conflicting_taxa":conflict,
        "wfo_taxa_with_invalid_name_or_missing_family":invalid,
        "cstd_identity":cstd_audit,
        "by_source":{},
        "original_external_protocol_state": "UNTESTED: no original-source independence or dual-tree matched-tip qualification",
    }
    for source in SOURCES:
        inp=getattr(a,source)
        # Single pinned commit; combined, cv and WFO also have SHA256 frozen.
        check(inp, SHA256.get(source))
        traits, rows, matched = source_presence(inp,f)
        stats=combined_support(cstd,traits,f,source)
        stats["source_rows"]=rows
        stats["source_rows_with_WFO_species"]=matched
        result["by_source"][source]=stats
    leaf_states,rows,unknown,accepted=cv_presence(a.cv,f)
    cv_report=combined_support(cstd,leaf_states,f,"herbarium_cv")
    cv_report["leaf_measurement_rows"]=rows
    cv_report["leaf_rows_without_unambiguous_WFO_taxon"]=unknown
    cv_report["leaf_rows_with_WFO_taxon"]=accepted
    result["by_source"]["herbarium_cv"]=cv_report
    result["source_family_diversity"]={
        "number_of_families_fruit_20plus_in_any_source":
            len(set().union(*[{f for f in entry["family_names"] if any(
                x["family"]==f and x["reproductive"].startswith("fruit_")
                for x in entry["pair_support"])}
                for entry in result["by_source"].values()])),
        "number_of_families_cv_independent_leaf_modality":
            cv_report["family_count"]
    }
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.out.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print("CSTD_NHM_ORGAN_SUPPORT="+json.dumps({
        "status":result["status"],"cstd_identity":cstd_audit,
        "by_source":{k:{
            x:v[x] for x in ("same_taxon_20plus_family_trait_pair_count",
                            "family_count","families_fruit","families_seed",
                            "reproductive_trait_names","vegetative_trait_names")
        } for k,v in result["by_source"].items()},
        "source_family_diversity":result["source_family_diversity"],
    },ensure_ascii=False))

if __name__=="__main__":
    main()
