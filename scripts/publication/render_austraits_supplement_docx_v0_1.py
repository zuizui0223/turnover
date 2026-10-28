#!/usr/bin/env python3
"""Render coauthor supplement as native Word landscape tables (reporting only).

Four tables are rendered directly, avoiding a broken Pandoc-generated table
style in this LibreOffice runtime. Values and source logic are untouched.
"""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WIDTHS={
 6:[2.15,1.65,.83,1.70,.85,2.92],
 7:[1.05,1.30,1.28,1.09,1.10,1.15,3.23],
 5:[.85,1.35,2.35,2.78,2.87],
 4:[.85,3.15,3.25,2.95],
}

def plain(s):
    s=re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)",r"\1 — \2",s)
    return s.replace("**","").replace("*","").replace(chr(96),"").strip()

def new_doc():
    d=Document()
    sec=d.sections[0]
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Cm(29.7);sec.page_height=Cm(21.0)
    sec.left_margin=sec.right_margin=Cm(1.65)
    sec.top_margin=sec.bottom_margin=Cm(1.6)
    normal=d.styles["Normal"]
    normal.font.name="Liberation Serif"
    normal.font.size=Pt(9.5)
    normal.paragraph_format.space_after=Pt(5)
    for h,size in [("Heading 1",15),("Heading 2",11.7),("Heading 3",10.4)]:
        st=d.styles[h];st.font.name="Liberation Sans"
        st.font.size=Pt(size);st.font.bold=True
        st.font.color.rgb=RGBColor(38,76,118)
        st.paragraph_format.space_before=Pt(9)
        st.paragraph_format.space_after=Pt(4)
    foot=sec.footer.paragraphs[0]
    foot.text="AusTraits — post-discovery study • Coauthor review only"
    foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in foot.runs:run.font.size=Pt(8)
    return d

def make_table(d,lines):
    rows=[]
    for line in lines:
        if re.fullmatch(r"\|\s*[-:|\s]+\|",line.strip()):
            continue
        rows.append([plain(c) for c in line.strip().strip("|").split("|")])
    if not rows:raise ValueError("empty supplementary table")
    ncol=len(rows[0])
    if ncol not in WIDTHS or any(len(r)!=ncol for r in rows):
        raise ValueError("supplement table column drift")
    tab=d.add_table(rows=len(rows),cols=ncol)
    tab.style="Table Grid";tab.autofit=False
    widths=WIDTHS[ncol]
    for j,w in enumerate(widths):tab.columns[j].width=Inches(w)
    trp=tab.rows[0]._tr.get_or_add_trPr()
    flag=OxmlElement("w:tblHeader");flag.set(qn("w:val"),"true")
    trp.append(flag)
    for i,values in enumerate(rows):
        for j,val in enumerate(values):
            cell=tab.rows[i].cells[j];cell.width=Inches(widths[j])
            cell.text=val
            if i==0:
                shd=OxmlElement("w:shd")
                shd.set(qn("w:fill"),"EEF2F5")
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1.3)
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.name="Liberation Sans"
                    run.font.size=Pt(7.6);run.bold=i==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return ncol,len(rows)

def build(source,out,manuscript,brier_figure):
    md=source.read_text(encoding="utf-8")
    for phrase in ("313/406 (77.1%)","314/397 (79.1%)",
                   "110/142 (77.5%)","120/142 (84.5%)",
                   "81/101 (80.2%)","non-nested","post-discovery"):
        if phrase.casefold() not in md.casefold():
            raise ValueError("Source table changed: "+phrase)
    d=new_doc()
    d.add_heading("Supplementary Tables — organ-function K",level=1)
    d.add_paragraph("Coauthor scientific review • 28 October 2026. "
       "All values were reported in existing post-discovery AusTraits analyses. "
       "Source/provenance and sampling caveats are unchanged.")
    lines=md.splitlines()
    i=0;specs=[]
    while i<len(lines):
        row=lines[i].strip()
        if not row:
            i+=1;continue
        if row.startswith("|"):
            block=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip());i+=1
            specs.append(make_table(d,block));continue
        if row.startswith("# "):
            i+=1;continue
        if row.startswith("## "):
            name=plain(row[3:]);p=d.add_heading(name,level=2)
            if name.startswith("Table S2.") or name.startswith("Table S3."):
                p.paragraph_format.page_break_before=True
            i+=1;continue
        if row.startswith("### "):
            d.add_heading(plain(row[4:]),level=3);i+=1;continue
        if row.startswith("- "):
            d.add_paragraph(plain(row[2:]),style="List Bullet");i+=1;continue
        buf=[]
        while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#","|","- ")):
            buf.append(plain(lines[i]));i+=1
        if buf:
            p=d.add_paragraph(" ".join(buf))
            p.paragraph_format.keep_together=True
        else:i+=1
    if specs!=[(6,6),(7,3),(5,3),(4,3)]:
        raise ValueError("Unexpected original 4-table schema: "+repr(specs))
    # Secondary evidence only: append the post-selected Brier figure to the
    # supplement rather than giving it main-manuscript prominence.
    src=manuscript.read_text(encoding="utf-8")
    marker="**Supplementary Figure S1.** "
    start=src.find(marker)
    if start<0:
        raise ValueError("Source supplemental figure caption missing")
    caption=src[start+len(marker):].split("\n\n",1)[0].strip()
    if "64.47%" not in caption or "non-nested" not in caption:
        raise ValueError("Supplementary caption does not warn against independent biological interpretation")
    if not brier_figure.exists() or brier_figure.stat().st_size<4000:
        raise ValueError("Source-linked Brier figure missing")
    fig_title=d.add_heading("Supplementary Figure S1. Conditional Brier scoring",level=2)
    fig_title.paragraph_format.page_break_before=True
    p=d.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run()
    shape=run.add_picture(str(brier_figure),width=Inches(8.0))
    shape._inline.docPr.set("descr",
        "Conditional cross-family Brier probability loss of pooled, organ-only, named-trait, and named-pair models; hard accuracy is algebraically redundant with original reproductive K wins.")
    cap=d.add_paragraph("Supplementary Figure S1. "+plain(caption))
    cap.paragraph_format.keep_together=True
    d.save(out)
    return specs

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--markdown",required=True,type=Path)
    p.add_argument("--output",required=True,type=Path)
    p.add_argument("--manuscript",required=True,type=Path)
    p.add_argument("--brier-figure",required=True,type=Path)
    a=p.parse_args()
    a.output.parent.mkdir(parents=True,exist_ok=True)
    print("NATIVE_WORD_SUPPLEMENT_TABLES_PASS",build(a.markdown,a.output,a.manuscript,a.brier_figure))

if __name__=="__main__":
    main()
