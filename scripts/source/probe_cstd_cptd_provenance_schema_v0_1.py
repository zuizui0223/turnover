#!/usr/bin/env python3
"""Pin and inspect only original-source/trait ontology METADATA of CSTD and CPTD.

This script never opens or summarizes numerical phenotype values, phylogenetic
trees, K, effect sizes or outcomes. It only identifies source-column schema,
site provenance labels and trait dictionary definitions; no PASS is inferred.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = {
    "cstd": "9d02dfed7de29484d2f5c5137235cabe",
    "sites": "76debd1ec32b850efa2c593ee5b94ee4",
}
TARGETS = {"plantheight","fruitlength","fruitwidth","seedlength","seedwidth",
           "leafarea","leafmassperarea","lma","averagela"}


def check(path, digest):
    h = hashlib.md5()
    with path.open("rb") as src:
        for b in iter(lambda: src.read(1024 * 1024), b""):
            h.update(b)
    if h.hexdigest() != digest:
        raise RuntimeError("Source file checksum mismatch, fail closed")


def norm(v):
    return re.sub("[^a-z0-9]", "", str(v or "").lower())


def sheet_metadata(ws, limit=5):
    rows = ws.iter_rows(values_only=True)
    first = next(rows, ())
    next_rows = []
    for i, r in enumerate(rows):
        if i >= limit:
            break
        next_rows.append([str(x).strip()[:180] if x is not None else ""
                          for x in r[:15]])
    return {
        "sheet": ws.title,
        "row_count_including_header": ws.max_row,
        "columns": [str(v or "")[:100] for v in first[:30]],
        "first_few_source_or_dictionary_metadata_rows": next_rows
    }


def run(a):
    check(a.cstd, EXPECTED["cstd"])
    check(a.sites, EXPECTED["sites"])
    wb = load_workbook(a.cstd, read_only=True, data_only=True)
    out = {
        "status": "SOURCE_METADATA_ONLY",
        "files_verified_md5": True,
        "no_phenotype_values_opened": True,
        "no_phylogenetic_K_computed": True,
        "cstd_metadata": {
            sh: sheet_metadata(wb[sh])
            for sh in ("Reference", "Data description", "Trait description")
        },
    }
    trait_dictionary = []
    for row in wb["Trait description"].iter_rows(min_row=2, values_only=True):
        label = str(row[0] or "").strip()
        if norm(label) in TARGETS:
            trait_dictionary.append({
                "trait": label,
                "verbatim_definition": str(row[1] or "").strip()[:800],
            })
    out["cstd_metadata"]["frozen_trait_definitions_only"] = trait_dictionary
    with a.sites.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        columns = reader.fieldnames or []
        original_source = next((x for x in columns if norm(x) in {
            "source", "reference", "datasource", "originsource"
        }), None)
        site = next((x for x in columns if norm(x)=="siteid"), None)
        count = Counter()
        n = 0
        examples = []
        for row in reader:
            n += 1
            if original_source:
                count[str(row.get(original_source) or "").strip()] += 1
            if len(examples) < 6:
                examples.append({
                    "site": (row.get(site, "") if site else "")[:80],
                    "source": (row.get(original_source, "") if original_source else "")[:180],
                })
    out["cptd_site_metadata"] = {
        "columns": columns, "rows": n,
        "source_field": original_source, "distinct_source_labels": len(count),
        "source_top_labels": count.most_common(30), "first_few_source_metadata_rows": examples,
    }
    return out


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--cstd", required=True, type=Path)
    p.add_argument("--sites", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    a = p.parse_args()
    out = run(a)
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("CSTD_CPTD_PROVENANCE_SCHEMA="+json.dumps(out,ensure_ascii=False))


if __name__ == "__main__":
    main()
