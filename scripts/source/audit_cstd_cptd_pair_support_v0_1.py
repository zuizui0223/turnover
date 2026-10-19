#!/usr/bin/env python3
"""Pre-outcome CSTD × CPTD-v2 original-taxon overlap audit.

Reads only accepted taxon IDs, family, original source labels, and presence
of each frozen trait. Never reads the magnitudes of leaf or reproductive
traits into statistics; no K, p values, effect selection or mechanism claim.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_MD5 = {
    "cstd": "9d02dfed7de29484d2f5c5137235cabe",
    "translation": "2c1285a6bf78112c855f0ce735685343",
    "chemical": "bf772a9e6a7251f8bed33d06d29f8fed",
}
CSTD_TRAITS = {
    "fruitlength": "fruit_length", "fruitwidth": "fruit_width",
    "seedlength": "seed_length", "seedwidth": "seed_width",
    "plantheight": "plant_height",
}
CPTD_TRAITS = {"Average LA": "leaf_area_candidate", "LMA": "leaf_mass_per_area_candidate"}
REPS = {"fruit_length", "fruit_width", "seed_length", "seed_width"}
VEGS = {"plant_height", "leaf_area_candidate", "leaf_mass_per_area_candidate"}
BAD = {"", "na", "n/a", "nan", "null", "none", "missing", "-", "n.d.", "nd"}


def verify(path, expected):
    h = hashlib.md5()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    if h.hexdigest() != expected:
        raise ValueError("Pinned file identity mismatch: " + str(path))


def tokenized_binomial(t):
    # Accepted names only. No fuzzy matches, genus-only assignments or taxon synonyms.
    parts = re.findall(r"[A-Za-z]+", str(t or ""))
    if len(parts) < 2 or parts[1].lower() in {"sp", "spp", "cf", "aff"}:
        return None
    return " ".join(p.lower() for p in parts[:2])


def from_translation(row):
    genus = str(row["ACCEPTED GENUS"] or "").strip()
    species = str(row["ACCEPTED SPECIES"] or "").strip()
    if not genus or not species:
        return None
    name = tokenized_binomial(species)
    if name:
        if name.split()[0] != genus.lower():
            return None
        return name
    # Source sometimes separates an accepted genus and one-word epithet.
    parts = re.findall(r"[A-Za-z]+", species)
    if len(parts) == 1 and parts[0].lower() not in {"sp", "spp"}:
        return tokenized_binomial(genus + " " + parts[0])
    return None


def present(v):
    # Presence only: don't parse, compare, transform or summarize numeric values.
    return v is not None and str(v).strip().lower() not in BAD


def audit(args):
    for kind, path in (("cstd", args.cstd), ("translation", args.translation), ("chemical", args.chemical)):
        verify(path, EXPECTED_MD5[kind])

    book = load_workbook(args.cstd, read_only=True, data_only=True)
    ws = book["Seed trait data"]
    head = next(ws.iter_rows(values_only=True))
    cols = {str(k): i for i, k in enumerate(head)}
    required = ("Species_accepted", "Family_accepted", "Trait", "Value", "Reference")
    if not all(k in cols for k in required):
        raise ValueError("CSTD original metadata schema changed")
    cstd = defaultdict(set)
    families_by_name = defaultdict(set)
    original_ref_count = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not present(row[cols["Value"]]):
            continue
        trait_label = re.sub(r"[^a-z0-9]", "", str(row[cols["Trait"]] or "").lower())
        trait = CSTD_TRAITS.get(trait_label)
        if not trait:
            continue
        species = tokenized_binomial(row[cols["Species_accepted"]])
        family = str(row[cols["Family_accepted"]] or "").strip()
        ref = str(row[cols["Reference"]] or "").strip()
        if not species or not family or not ref:
            continue
        cstd[(family, species)].add(trait)
        families_by_name[species].add(family)
        original_ref_count[ref] += 1

    # A binomial mapped to conflicting CSTD family names is excluded rather than
    # forcing a taxonomic identity or accidentally duplicating support.
    conflicted = {k for k, v in families_by_name.items() if len(v) != 1}
    species_to_family = {k: next(iter(v)) for k, v in families_by_name.items()
                         if len(v) == 1}

    with args.translation.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        required = {"SAMPLE ID", "ACCEPTED GENUS", "ACCEPTED SPECIES"}
        if not required.issubset(set(reader.fieldnames or [])):
            raise ValueError("CPTD species-translation metadata schema changed")
        samples = {}
        for row in reader:
            name = from_translation(row)
            if name:
                samples[str(row["SAMPLE ID"]).strip()] = name

    cptd = defaultdict(set)
    taxon_samples = Counter()
    with args.chemical.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        required = {"SAMPLE ID", "Average LA", "LMA"}
        if not required.issubset(set(reader.fieldnames or [])):
            raise ValueError("CPTD leaf-chemical metadata schema changed")
        for row in reader:
            taxon = samples.get(str(row["SAMPLE ID"]).strip())
            if taxon is None:
                continue
            taxon_samples[taxon] += 1
            for col, trait in CPTD_TRAITS.items():
                if present(row[col]):
                    cptd[taxon].add(trait)

    merged = defaultdict(set)
    for (family, species), traits in cstd.items():
        if species in conflicted:
            continue
        merged[(family, species)].update(traits)
    for species, traits in cptd.items():
        family = species_to_family.get(species)
        if family is not None:
            merged[(family, species)].update(traits)

    by_family_trait = defaultdict(lambda: defaultdict(set))
    for (family, species), traits in merged.items():
        for trait in traits:
            by_family_trait[family][trait].add(species)
    pairs = []
    for family, by_trait in sorted(by_family_trait.items()):
        for rep in sorted(REPS & set(by_trait)):
            for veg in sorted(VEGS & set(by_trait)):
                n = len(by_trait[rep] & by_trait[veg])
                if n >= 20:
                    pairs.append({"family": family, "reproductive": rep, "vegetative": veg,
                                  "same_accepted_binomial": n, "vegetative_source":
                                  "CSTD" if veg == "plant_height" else "CPTD_v2"})
    matched_fam = {x["family"] for x in pairs}
    fruit = {x["family"] for x in pairs if x["reproductive"].startswith("fruit_")}
    seed = {x["family"] for x in pairs if x["reproductive"].startswith("seed_")}
    veg = {x["vegetative"] for x in pairs}
    rep = {x["reproductive"] for x in pairs}
    cross = [x for x in pairs if x["vegetative_source"] == "CPTD_v2"]

    return {
        "status": "HOLD_PRE_OUTCOME",
        "dataset_primary_sources": [
            "CSTD v4 10.6084/m9.figshare.28152035.v4 (compiled 694 sources)",
            "CPTD v2 10.6084/m9.figshare.19448219 (mixed original-field and literature sites)"
        ],
        "source_independence": "UNKNOWN; source DOI/record-level reuse and original sampling protocols not verified",
        "matched_tree_eligibility": "UNTESTED: no source-native S3 or prune crosswalk",
        "trait_homology": "UNVERIFIED: CPTD Average LA and LMA are candidates, not approved AusTraits ontology equivalents",
        "observed_magnitudes_or_phylogenetic_effects_opened": False,
        "source_identity_md5_verified": True,
        "metadata_support": {
            "cstd_accepted_binomial_count": len(species_to_family),
            "cstd_conflicting_binomials_excluded": len(conflicted),
            "cptd_accepted_binomial_count": len(taxon_samples),
            "exact_binomial_taxon_overlap": len(set(species_to_family) & set(taxon_samples)),
            "species_with_CPTD_leaf_metadata_and_CSTD_reproductive_labels":
                sum(any(t in REPS for t in traits) and bool(cptd.get(sp) & set(CPTD_TRAITS.values()))
                    for (_fa, sp), traits in cstd.items() if sp not in conflicted),
            "same_binomial_20plus_pairs_including_CSTD_height": len(pairs),
            "families_with_any_pair": len(matched_fam),
            "families_fruit": len(fruit),
            "families_seed": len(seed),
            "vegetative_trait_candidate_names": sorted(veg),
            "reproductive_trait_candidate_names": sorted(rep),
            "cross_source_20plus_pairs": len(cross),
            "cross_source_20plus_families": len({x["family"] for x in cross}),
            "cross_source_pairs": cross,
        },
        "all_necessary_support_dimensions_possible_unverified":
            len(matched_fam) >= 12 and len(fruit) >= 6 and len(seed) >= 6
            and len(rep) >= 2 and len(veg) >= 2,
        "explicit_hard_stop": "Even positive necessary support cannot authorize K. Source independence, original-study lineage, numeric morphology concepts, permissions and tree support are NOT checked."
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--cstd", type=Path, required=True)
    p.add_argument("--translation", type=Path, required=True)
    p.add_argument("--chemical", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    a = p.parse_args()
    result = audit(a)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    a.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("CSTD_CPTD_SOURCE_ONLY=" + json.dumps({
        "status": result["status"], **{k: v for k, v in result["metadata_support"].items()
        if k != "cross_source_pairs"},
        "necessary_gate_possible_unverified": result["all_necessary_support_dimensions_possible_unverified"],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
